# -*- coding: utf-8 -*-
"""Gộp điểm Hóa Lần 3 (12A10) vào data.json.
Idempotent: chạy lại sau merge_update.py để phục hồi lớp L3.
Không đụng số liệu L1/L2 toàn khối (L3 mới chỉ có Hóa 12A10).
"""
import json
import os
import re
from datetime import datetime
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(ROOT, "data.json")
L3FILE = r"C:\Users\LENOVO\Downloads\12A10- KTĐK-LẦN 3.xlsx"
MON = "Hóa học"


def norm(s):
    return re.sub(r"\s+", " ", str(s).strip()) if s not in (None, "") else ""


def avg(xs):
    xs = [x for x in xs if x is not None]
    return round(sum(xs) / len(xs), 3) if xs else None


data = json.load(open(DATA, encoding="utf-8"))
bykey = {(s["lop"], norm(s["ten"])): s for s in data["students"]}

wb = load_workbook(L3FILE, data_only=True, read_only=True)
ws = wb["Lần 2"]
l3map, lạ = {}, 0
for r in list(ws.iter_rows(values_only=True))[4:]:
    if r[0] in (None, ""):
        continue
    ten = norm("{} {}".format(r[1], r[2]))
    try:
        v = float(r[4]) if r[4] not in (None, "") else None
    except (TypeError, ValueError):
        v = None
    if ("12A10", ten) not in bykey:
        lạ += 1
        print("LA:", repr(ten), v)
        continue
    l3map[ten] = v

n_set = 0
for ten, v in l3map.items():
    s = bykey[("12A10", ten)]
    s["l3"] = {MON: v}
    l2 = (s.get("l2") or {}).get(MON)
    s["delta23"] = {MON: round(v - l2, 2) if (v is not None and l2 is not None) else None}
    n_set += 1

vals = [v for v in l3map.values() if v is not None]
pairs = [(t, v) for t, v in l3map.items()
         if v is not None and (bykey[("12A10", t)].get("l2") or {}).get(MON) is not None]
st = data["class_stats"]["12A10"][MON]
st["n3"] = len(vals)
st["tb3"] = avg(vals)
st["duoi5_l3"] = sum(1 for x in vals if x < 5)

for grp in ("phudao", "quantam", "boiduong", "tienbo", "thutlui"):
    for item in data["class_lists"]["12A10"][MON][grp]:
        ten = norm(item["ten"])
        v = l3map.get(ten)
        item["l3"] = v
        l2 = item.get("l2")
        item["delta23"] = round(v - l2, 2) if (v is not None and l2 is not None) else None

meta = data["meta"]
meta["lan3"] = "KTĐK Lần 3 (môn Hóa, lớp 12A10, 13/09/2026)"
meta["cap_nhat"] = datetime.now().strftime("%H:%M %d/%m/%Y") + " + Hóa L3 12A10"
if os.path.basename(L3FILE) not in meta["nguon"]:
    meta["nguon"].append(os.path.basename(L3FILE))

json.dump(data, open(DATA, "w", encoding="utf-8"), ensure_ascii=False)
print("khớp L3: %d/23, lạ: %d" % (n_set, lạ))
print("Hóa 12A10 L3: n=%d TB=%s <5=%d" % (st["n3"], st["tb3"], st["duoi5_l3"]))
print("PD L3 (<5):", [(t, v) for t, v in sorted(l3map.items(), key=lambda x: (x[1] is None, x[1])) if v is not None and v < 5])
print("wrote", DATA, os.path.getsize(DATA))
