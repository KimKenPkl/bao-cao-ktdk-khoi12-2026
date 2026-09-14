# -*- coding: utf-8 -*-
"""Cap nhat diem KTĐK Lan 3 day du Toan-Hoa-Sinh lop 12A10 vao data.json.
Nguon chuan: 12A10_ DIEM KTDK Lan 3.xlsx (sheet 12A10-L3) - 23 em SBD 1001-1023.
Giu nguyen L1/L2 toan khoi. Chi them l3/delta23 + class_stats n3/tb3/duoi5_l3.
Idempotent: chay lai cho ket qua giong nhau.
"""
import json, os, re, shutil
from datetime import datetime
from openpyxl import load_workbook

ROOT = r'D:\bao-cao-ktdk-khoi12'
DATA = os.path.join(ROOT, 'data.json')
SRC = r'C:\Users\LENOVO\Downloads\12A10_ ĐIỂM KTĐK Lần 3.xlsx'
SHEET = '12A10-L3'
MONS = ['Toán', 'Hóa học', 'Sinh học']

def norm(s):
    return re.sub(r'\s+', ' ', str(s).strip()) if s not in (None, '') else ''

def fnum(v):
    try:
        return float(v) if v not in (None, '') else None
    except (TypeError, ValueError):
        return None

def avg(xs):
    xs = [x for x in xs if x is not None]
    return round(sum(xs)/len(xs), 3) if xs else None

# backup
bk = DATA + '.bak_%s' % datetime.now().strftime('%Y%m%d_%H%M%S')
shutil.copy2(DATA, bk)
print('backup:', bk)

data = json.load(open(DATA, encoding='utf-8'))
bykey = {(s['lop'], norm(s['ten'])): s for s in data['students']}

wb = load_workbook(SRC, data_only=True, read_only=True)
ws = wb[SHEET]
rows = list(ws.iter_rows(values_only=True))
print('rows:', len(rows), 'header:', rows[2])

# doc 23 dong HS (R4-R26)
l3_toan, l3_hoa, l3_sinh = {}, {}, {}
for r in rows[3:26]:
    stt, sbd, ho, ten, toan, hoa, sinh, van, b00 = (list(r) + [None]*9)[:9]
    if sbd in (None, ''):
        continue
    full = norm(f'{ho} {ten}')
    l3_toan[full] = fnum(toan)
    l3_hoa[full] = fnum(hoa)
    l3_sinh[full] = fnum(sinh)

print('doc duoc:', len(l3_toan), 'em')
# doi chieu TB voi dong tong trong file
tb_file = [fnum(c) for c in rows[28][4:7]]
print('TB trong file (Toan,Hoa,Sinh):', tb_file)
print('TB tinh lai:', avg(list(l3_toan.values())), avg(list(l3_hoa.values())), avg(list(l3_sinh.values())))

# khop ten voi web
la = []
for ten in l3_toan:
    if ('12A10', ten) not in bykey:
        la.append(ten)
print('la (khong khop web):', la)
missing = [k[1] for k in bykey if k[0]=='12A10' and k[1] not in l3_toan]
print('thieu (web co, file khong):', missing)
assert len(l3_toan)==23 and not la and not missing, 'KHOP TEN THAT BAI'

l3map = {'Toán': l3_toan, 'Hóa học': l3_hoa, 'Sinh học': l3_sinh}

# 1. students
n_set = 0
for ten in l3_toan:
    s = bykey[('12A10', ten)]
    if 'l3' not in s or not isinstance(s['l3'], dict):
        s['l3'] = {}
    if 'delta23' not in s or not isinstance(s['delta23'], dict):
        s['delta23'] = {}
    for mon in MONS:
        v = l3map[mon][ten]
        s['l3'][mon] = v
        l2 = (s.get('l2') or {}).get(mon)
        s['delta23'][mon] = round(v-l2, 2) if (v is not None and l2 is not None) else None
    n_set += 1
print('cap nhat students:', n_set)

# 2. class_stats
for mon in MONS:
    vals = [v for v in l3map[mon].values() if v is not None]
    st = data['class_stats']['12A10'][mon]
    st['n3'] = len(vals)
    st['tb3'] = avg(vals)
    st['duoi5_l3'] = sum(1 for x in vals if x < 5)
    print(f"{mon}: n3={st['n3']} tb3={st['tb3']} <5={st['duoi5_l3']} (cu: n1={st['n1']} tb1={st['tb1']} | n2={st['n2']} tb2={st['tb2']})")

# 3. class_lists (giu nhom L2, chi them l3/delta23)
for mon in MONS:
    for grp in ('phudao','quantam','boiduong','tienbo','thutlui'):
        for item in data['class_lists']['12A10'][mon][grp]:
            ten = norm(item['ten'])
            v = l3map[mon].get(ten)
            item['l3'] = v
            l2 = item.get('l2')
            item['delta23'] = round(v-l2, 2) if (v is not None and l2 is not None) else None
print('cap nhat class_lists xong')

# 4. meta
meta = data['meta']
meta['lan3'] = 'KTĐK Lần 3 (Toán–Hóa–Sinh, lớp 12A10, 14/09/2026)'
meta['cap_nhat'] = datetime.now().strftime('%H:%M %d/%m/%Y') + ' + L3 full 12A10'
base = os.path.basename(SRC)
if base not in meta['nguon']:
    meta['nguon'].append(base)
print('meta lan3:', meta['lan3'])

json.dump(data, open(DATA, 'w', encoding='utf-8'), ensure_ascii=False)
print('wrote', DATA, os.path.getsize(DATA))

# tom tat nhanh de kiem tra
vals_T = sorted(l3_toan.items(), key=lambda x: x[1])
vals_H = sorted(l3_hoa.items(), key=lambda x: x[1])
vals_S = sorted(l3_sinh.items(), key=lambda x: x[1])
print('Toan <5:', [(t,v) for t,v in vals_T if v<5])
print('Hoa <5:', [(t,v) for t,v in vals_H if v<5])
print('Sinh <5:', [(t,v) for t,v in vals_S if v<5][:8], '... tong', sum(1 for _,v in vals_S if v<5))
