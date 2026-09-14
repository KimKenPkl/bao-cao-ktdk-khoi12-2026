# -*- coding: utf-8 -*-
"""Gop diem KTĐK Lan 3 TOAN KHOI (11 lop) tu file tong vao data.json web.
Nguon: 2026-2027_HK1_Khoi12_Diem-KTDK-Lan_3.xlsx (da sua link Hoa/12A4).
- students: them s['l3'][mon], s['delta23'][mon] (khop ten + fuzzy + them HS moi).
- class_stats + overview_classes: n3/tb3/duoi5_l3 theo file.
- class_lists: item['l3']/item['delta23'].
- meta: lan3/cap_nhat/nguon.
Idempotent khi chay lai.
"""
import json, os, re, shutil
from datetime import datetime
from openpyxl import load_workbook

ROOT = r'D:\bao-cao-ktdk-khoi12'
DATA = os.path.join(ROOT, 'data.json')
SRC = r'C:\Users\LENOVO\Downloads\2026-2027_HK1_Khoi12_Diem-KTDK-Lan_3.xlsx'
LOPS = ['12A1', '12A2', '12A3', '12A4', '12A5', '12A6',
        '12A7', '12A8', '12A9', '12A10', '12A11']
SUBJECTS = ['Toán', 'Ngữ văn', 'Vật lí', 'Hóa học', 'Sinh học',
            'Lịch sử', 'Tin học', 'Anh Văn', 'KTPL']
HDR2MON = {'toán': 'Toán', 'lý': 'Vật lí', 'hóa': 'Hóa học', 'văn': 'Ngữ văn',
           'anh': 'Anh Văn', 'sinh': 'Sinh học', 'sử': 'Lịch sử',
           'tin': 'Tin học', 'ktpl': 'KTPL'}

def norm(s):
    return re.sub(r'\s+', ' ', str(s).strip()) if s not in (None, '') else ''

def fnum(v):
    try:
        return float(v) if v not in (None, '') else None
    except (TypeError, ValueError):
        return None

def avg(xs):
    xs = [x for x in xs if x is not None]
    return round(sum(xs) / len(xs), 3) if xs else None

def dedup_tok(s):
    out = []
    for t in norm(s).split(' '):
        if not out or out[-1] != t:
            out.append(t)
    return ' '.join(out)

bk = DATA + '.bak_%s' % datetime.now().strftime('%Y%m%d_%H%M%S')
shutil.copy2(DATA, bk)
print('backup:', bk)

data = json.load(open(DATA, encoding='utf-8'))
bykey = {(s['lop'], norm(s['ten'])): s for s in data['students']}

wb = load_workbook(SRC, data_only=True, read_only=True)
n_match = n_fuzzy = n_new = 0
fuzzy_log, new_log = [], []
file_vals = {}  # (lop, mon) -> [scores>0]

for lop in LOPS:
    ws = wb[lop]
    hmap = {}
    for c in range(1, ws.max_column + 1):
        h = norm(ws.cell(3, c).value).lower()
        for key, mon in HDR2MON.items():
            if h.startswith(key):
                hmap[c] = mon
                break
    for mon in hmap.values():
        file_vals.setdefault((lop, mon), [])
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, 2).value in (None, ''):
            break
        full = norm(f"{ws.cell(r, 3).value or ''} {ws.cell(r, 4).value or ''}")
        if not full:
            continue
        scores = {mon: fnum(ws.cell(r, c).value) for c, mon in hmap.items()}
        key = (lop, full)
        rec = bykey.get(key)
        if rec is None:
            cands = [k for k in bykey if k[0] == lop and (
                dedup_tok(k[1]) == dedup_tok(full) or k[1].startswith(full)
                or full.startswith(k[1]))]
            if len(cands) == 1:
                rec = bykey[cands[0]]
                n_fuzzy += 1
                fuzzy_log.append(f'{lop}: file[{full}] -> web[{cands[0][1]}]')
            else:
                rec = {'lop': lop, 'ten': full,
                       'l1': {m: None for m in SUBJECTS},
                       'l2': {m: None for m in SUBJECTS},
                       'delta': {m: None for m in SUBJECTS}}
                data['students'].append(rec)
                bykey[key] = rec
                n_new += 1
                new_log.append(f'{lop}: +{full}')
        else:
            n_match += 1
        if 'l3' not in rec or not isinstance(rec['l3'], dict):
            rec['l3'] = {}
        if 'delta23' not in rec or not isinstance(rec['delta23'], dict):
            rec['delta23'] = {}
        for mon, v in scores.items():
            rec['l3'][mon] = v
            l2 = (rec.get('l2') or {}).get(mon)
            rec['delta23'][mon] = round(v - l2, 2) if (v is not None and l2 is not None) else None
            if v is not None and v > 0:
                file_vals[(lop, mon)].append(v)

print(f'match={n_match} fuzzy={n_fuzzy} new={n_new}')
for x in fuzzy_log:
    print('  FUZZY', x)
for x in new_log:
    print('  NEW', x)

# class_stats + overview_classes
ov_by_lop = {o['lop']: o for o in data['overview_classes']}
for lop in LOPS:
    for mon in SUBJECTS:
        vals = file_vals.get((lop, mon), [])
        st = data['class_stats'].setdefault(lop, {}).setdefault(mon, {})
        st['n3'] = len(vals)
        st['tb3'] = avg(vals)
        st['duoi5_l3'] = sum(1 for x in vals if x < 5)
        if lop in ov_by_lop:
            ov = ov_by_lop[lop]['stats'].setdefault(mon, {})
            ov['n3'] = st['n3']
            ov['tb3'] = st['tb3']
            ov['duoi5_l3'] = st['duoi5_l3']
    tot = sum(len(file_vals.get((lop, m), [])) for m in SUBJECTS)
    print(f'{lop}: tong luot thi L3={tot} | Toan n3={data["class_stats"][lop]["Toán"]["n3"]} '
          f'tb3={data["class_stats"][lop]["Toán"]["tb3"]}')

# class_lists
for lop in LOPS:
    cl = data['class_lists'].get(lop, {})
    for mon, groups in cl.items():
        for grp, items in groups.items():
            for item in items:
                ten = norm(item.get('ten', ''))
                rec = bykey.get((lop, ten))
                v = (rec.get('l3') or {}).get(mon) if rec else None
                item['l3'] = v
                l2 = item.get('l2')
                item['delta23'] = round(v - l2, 2) if (v is not None and l2 is not None) else None

# meta
meta = data['meta']
meta['lan3'] = 'KTĐK Lần 3 (11 lớp khối 12, 14/09/2026 — đã sửa link Hóa/12A4)'
meta['cap_nhat'] = datetime.now().strftime('%H:%M %d/%m/%Y') + ' + L3 full 11 lớp'
base = os.path.basename(SRC)
if base not in meta['nguon']:
    meta['nguon'].append(base)

json.dump(data, open(DATA, 'w', encoding='utf-8'), ensure_ascii=False)
print('wrote', DATA, os.path.getsize(DATA), '| tong HS:', len(data['students']))
