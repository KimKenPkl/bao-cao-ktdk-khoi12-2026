# -*- coding: utf-8 -*-
"""Gộp batch file lớp 11/09 (Văn L1 nhiều lớp + 12A6 L2 + hiệu chỉnh) vào data.json."""
import openpyxl, glob, os, json, re
from datetime import datetime

D = r"C:\Users\LENOVO\Downloads"
OUT = r"D:\bao-cao-ktdk-khoi12\data.json"
data = json.load(open(OUT, encoding="utf-8"))

SUBJECTS = ["Toán","Ngữ văn","Vật lí","Hóa học","Sinh học","Lịch sử","Tin học","Anh Văn","KTPL"]

def norm(s):
    return re.sub(r"\s+", " ", str(s).strip()) if s not in (None, "") else ""

def fnum(v):
    try: return float(v) if v not in (None, "") else None
    except: return None

# dựng dict base từ data.json hiện tại
d1 = {(s["lop"], s["ten"]): dict(s["l1"]) for s in data["students"]}
d2 = {(s["lop"], s["ten"]): dict(s["l2"]) for s in data["students"]}
assert len(d1) >= 450 and len(d2) >= 450, f"BASE LỖI: L1={len(d1)} L2={len(d2)}"
assert ("12A7", "Vũ Đoàn Duy Anh") in d1, "BASE LỖI: thiếu key mẫu"
print(f"BASE OK: L1={len(d1)} L2={len(d2)}")
log = []
STAT = {"khop": 0, "them_diem": 0, "sua": 0, "hs_moi": 0}

def upsert(lop, ten, rnd, smap):
    """smap: {mon: value}. Thêm HS mới nếu chưa có."""
    global d1, d2
    DD = d1 if rnd == "L1" else d2
    key = (lop, ten)
    if key not in DD:
        DD[key] = {}
        STAT["hs_moi"] += 1
        log.append(f"  + HS mới {lop} {ten} ({rnd})")
    else:
        STAT["khop"] += 1
    for s, v in smap.items():
        nv = fnum(v)
        if nv is None: continue
        ov = DD[key].get(s)
        if ov is None:
            DD[key][s] = nv; STAT["them_diem"] += 1
        elif abs(ov - nv) > 1e-9:
            log.append(f"  ~ {lop} {ten} {rnd} {s}: {ov} -> {nv}")
            DD[key][s] = nv; STAT["sua"] += 1

def read_simple(path, sheet, lop, rnd, scol, ho_ten_cols, header_row=4, start_row=5):
    """File 1 khối: header_row (1-indexed) chứa tên cột; scol={mon: header_name}."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [norm(c) for c in rows[header_row-1]]
    def col(h):
        return hdr.index(h) if h in hdr else None
    iho, iten = ho_ten_cols
    if isinstance(iho, str): iho = col(iho)
    if isinstance(iten, str): iten = col(iten)
    n = 0
    for r in rows[start_row-1:]:
        if len(r) <= max([c for c in [iho, iten] if c is not None]+[0]): continue
        ho = norm(r[iho]) if iho is not None else ""
        tn = norm(r[iten]) if iten is not None else ""
        if not ho and not tn: continue
        if ho in ("HỌ VÀ TÊN", "HỌ", "Họ") or tn in ("Tên",): continue
        if "ĐIỂM" in ho.upper() or "ĐIỂM" in tn.upper() or "DƯỚI" in ho.upper(): continue
        ten = norm(f"{ho} {tn}")
        if not ten or ten.lower() == "none none": continue
        smap = {}
        for mon, hname in scol.items():
            ci = col(hname)
            if ci is not None and len(r) > ci:
                smap[mon] = r[ci]
        if not any(fnum(v) is not None for v in smap.values()): continue
        upsert(lop, ten, rnd, smap); n += 1
    msg = f"{os.path.basename(path)}[{sheet}] {lop}-{rnd}: {n} dòng"
    print(msg); log.append(msg)

M = "2026-09-11 22:00"
import time
cut = time.mktime(time.strptime(M, "%Y-%m-%d %H:%M"))
newfiles = sorted([f for f in glob.glob(os.path.join(D, "*.xlsx")) if os.path.getmtime(f) > cut],
                  key=os.path.basename)
print("FILES:", [os.path.basename(f) for f in newfiles])

# 1. 12A1 LẦN 1: SBD,HỌ,TÊN,TOÁN,LÝ,HÓA,AV,VĂN
read_simple(os.path.join(D, "ĐIỂM KTĐK 12A1 LẦN 1.xlsx"), "LẦN 1", "12A1", "L1",
            {"Toán":"TOÁN","Vật lí":"LÝ","Hóa học":"HÓA","Anh Văn":"AV","Ngữ văn":"VĂN"}, ("HỌ","TÊN"))
# 2. 12A2 TOÁN L1: SBD,HỌ VÀ TÊN(x2),Toán,Lý,AV,Văn
read_simple(os.path.join(D, "12A2_BẢNG ĐIỂM KTĐK TOÁN L1.xlsx"), "L1", "12A2", "L1",
            {"Toán":"Toán","Vật lí":"Lý","Anh Văn":"AV","Ngữ văn":"Văn"}, (2,3), header_row=4, start_row=5)
# 3. 12A3 LẦN 01: SMAS,SBD,Họ,Tên,Toán,Lí,Hóa,Văn
read_simple(os.path.join(D, "12A3 - ĐIỂM KTĐK LẦN 01 2026-2027.xlsx"), "LẦN 01", "12A3", "L1",
            {"Toán":"Toán","Vật lí":"Lí","Hóa học":"Hóa","Ngữ văn":"Văn"}, ("Họ","Tên"))
# 4. 12A5 Lần 1: SBD,Họ,Tên,Toán,Hóa,Sinh,Sinh,Anh,Sử,Văn (lấy theo vị trí)
p5 = os.path.join(D, "12A5 - Điểm KTĐK Lần 1.xlsx")
wb = openpyxl.load_workbook(p5, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]; rows = list(ws.iter_rows(values_only=True))
n5 = 0
for r in rows[4:]:
    if r[0] in (None, ""): continue
    ten = norm(f"{norm(r[1])} {norm(r[2])}")
    if not ten: continue
    smap = {"Toán": r[3], "Hóa học": r[4], "Sinh học": r[5], "Ngữ văn": r[9]}
    if not any(fnum(v) is not None for v in smap.values()): continue
    upsert("12A5", ten, "L1", smap); n5 += 1
print(f"12A5-L1: {n5} dòng"); log.append(f"12A5-L1: {n5} dòng")
# 5. 12A6 L1 + L2(1): SBD,HỌ,TÊN,TOÁN,VĂN,ANH,SỬ,KTPL
read_simple(os.path.join(D, "12A6_BẢNG ĐIỂM KTĐK_LẦN 1.xlsx"), "L1", "12A6", "L1",
            {"Toán":"TOÁN","Ngữ văn":"VĂN","Anh Văn":"ANH","Lịch sử":"SỬ","KTPL":"KTPL"}, ("HỌ","TÊN"))
read_simple(os.path.join(D, "12A6_BẢNG ĐIỂM KTĐK_LẦN 2 (1).xlsx"), "L2", "12A6", "L2",
            {"Toán":"TOÁN","Ngữ văn":"VĂN","Anh Văn":"ANH","Lịch sử":"SỬ","KTPL":"KTPL"}, ("HỌ","TÊN"))
# 6. 12A8 LẦN 1: SBD,HỌ VÀ TÊN(x2),Toán,Lý,AV,Văn
read_simple(os.path.join(D, "12A8_ĐIỂM  KTĐK LẦN 1.xlsx"), "LẦN 1", "12A8", "L1",
            {"Toán":"Toán","Vật lí":"Lý","Anh Văn":"AV","Ngữ văn":"Văn"}, (2,3))
# 7. 12A11: sheets Lần 2 + Lần 1: SBD,HỌ VÀ TÊN(x2),Toán,Lý,Hoá,Văn
read_simple(os.path.join(D, "12A11 - ĐIỂM KTĐK.xlsx"), "Lần 1", "12A11", "L1",
            {"Toán":"Toán","Vật lí":"Lý","Hóa học":"Hoá","Ngữ văn":"Văn"}, (2,3), header_row=3, start_row=4)
read_simple(os.path.join(D, "12A11 - ĐIỂM KTĐK.xlsx"), "Lần 2", "12A11", "L2",
            {"Toán":"Toán","Vật lí":"Lý","Hóa học":"Hoá","Ngữ văn":"Văn"}, (2,3), header_row=3, start_row=4)
# 8. 12A9 (1): sheets 'L 1' (Toán,Lý,Hóa) + 'L 2' (Toán,Lý,Hóa,Văn)
read_simple(os.path.join(D, "12A9_ ĐIỂM KTĐK Lần 2 (1).xlsx"), "L 1", "12A9", "L1",
            {"Toán":"Toán","Vật lí":"Lý","Hóa học":"Hóa"}, ("Họ","Tên"), header_row=3, start_row=4)
read_simple(os.path.join(D, "12A9_ ĐIỂM KTĐK Lần 2 (1).xlsx"), "L 2", "12A9", "L2",
            {"Toán":"Toán","Vật lí":"Lý","Hóa học":"Hóa","Ngữ văn":"Văn"}, ("Họ","Tên"), header_row=3, start_row=4)
# 9. 12A4 +VĂN: 2 khối (trái A1: Toán,Lý,AV,Văn / phải Tin: Toán,Tin,Lý,Văn), header row4, data row5+
p4 = os.path.join(D, "12A4 - ĐIỂM KTĐK LẦN 1 +VĂN.xlsx")
wb = openpyxl.load_workbook(p4, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]; rows = list(ws.iter_rows(values_only=True))
n4 = 0
for r in rows[4:]:
    # khối trái
    if r[0] not in (None, "") and norm(r[1]):
        ten = norm(f"{norm(r[1])} {norm(r[2])}")
        if ten:
            smap = {"Toán": r[3], "Vật lí": r[4], "Anh Văn": r[5], "Ngữ văn": r[6]}
            if any(fnum(v) is not None for v in smap.values()):
                upsert("12A4", ten, "L1", smap); n4 += 1
    # khối phải (cột 10-17)
    if len(r) > 11 and r[10] not in (None, "") and norm(r[11]):
        ten = norm(f"{norm(r[11])} {norm(r[12])}")
        if ten:
            smap = {"Toán": r[13], "Tin học": r[14], "Vật lí": r[15], "Ngữ văn": r[16]}
            if any(fnum(v) is not None for v in smap.values()):
                upsert("12A4", ten, "L1", smap); n4 += 1
print(f"12A4-L1(+VAN): {n4} dòng"); log.append(f"12A4-L1(+VAN): {n4} dòng")
# 10. 12A7 (3): kiểm tra hiệu chỉnh thêm
read_simple(os.path.join(D, "12A7_L2 (3).xlsx"), "L1", "12A7", "L1",
            {"Toán":"Toán","Vật lí":"Lý","Hóa học":"Hóa","Ngữ văn":"Văn"}, ("Họ","Tên"), header_row=3, start_row=4)
read_simple(os.path.join(D, "12A7_L2 (3).xlsx"), "L2", "12A7", "L2",
            {"Toán":"Toán","Vật lí":"Lý","Hóa học":"Hóa","Ngữ văn":"Văn"}, ("Họ","Tên"), header_row=3, start_row=4)

print("STAT:", STAT)

# ---------- tái tính ----------
all_keys = sorted(set(d1) | set(d2))
classes = sorted({k[0] for k in all_keys})
def avg(xs):
    xs = [x for x in xs if x is not None]
    return round(sum(xs)/len(xs), 3) if xs else None
def pct_below(xs, th=5):
    xs = [x for x in xs if x is not None]
    return round(100*sum(1 for x in xs if x < th)/len(xs), 1) if xs else None

students = []
for lop, ten in all_keys:
    rec = {"lop": lop, "ten": ten, "l1": d1.get((lop, ten), {}), "l2": d2.get((lop, ten), {})}
    rec["delta"] = {s: (round(rec["l2"].get(s)-rec["l1"].get(s), 2)
                        if (rec["l1"].get(s) is not None and rec["l2"].get(s) is not None) else None)
                    for s in SUBJECTS}
    students.append(rec)

subjects_stats = {}
for s in SUBJECTS:
    v1 = [d1[k].get(s) for k in d1 if d1[k].get(s) is not None]
    v2 = [d2[k].get(s) for k in d2 if d2[k].get(s) is not None]
    pd = [d2[k][s]-d1[k][s] for k in set(d1) & set(d2) if d1[k].get(s) is not None and d2[k].get(s) is not None]
    subjects_stats[s] = {"n1": len(v1), "n2": len(v2), "tb1": avg(v1), "tb2": avg(v2),
        "delta": round(avg(v2)-avg(v1), 3) if (avg(v1) is not None and avg(v2) is not None) else None,
        "duoi5_l1": sum(1 for x in v1 if x < 5), "duoi5_l2": sum(1 for x in v2 if x < 5),
        "pct_duoi5_l1": pct_below(v1), "pct_duoi5_l2": pct_below(v2),
        "gioi_l2": sum(1 for x in v2 if x >= 8), "kha_l2": sum(1 for x in v2 if 6.5 <= x < 8),
        "tb_y_l2": sum(1 for x in v2 if 5 <= x < 6.5),
        "tang": sum(1 for x in pd if x > 0), "giam": sum(1 for x in pd if x < 0), "bang": sum(1 for x in pd if x == 0)}

class_stats = {}; class_lists = {}
for lop in classes:
    class_stats[lop] = {}; class_lists[lop] = {}
    keys = [k for k in all_keys if k[0] == lop]
    for s in SUBJECTS:
        v1 = [d1[k].get(s) for k in keys if k in d1 and d1[k].get(s) is not None]
        v2 = [d2[k].get(s) for k in keys if k in d2 and d2[k].get(s) is not None]
        class_stats[lop][s] = {"n1": len(v1), "n2": len(v2), "tb1": avg(v1), "tb2": avg(v2),
            "delta": (round(avg(v2)-avg(v1), 2) if v1 and v2 else None),
            "duoi5_l1": sum(1 for x in v1 if x < 5), "duoi5_l2": sum(1 for x in v2 if x < 5)}
        ph, qu, bo, ti, th = [], [], [], [], []
        for k in keys:
            a = d1.get(k, {}).get(s); b = d2.get(k, {}).get(s)
            if b is None: continue
            item = {"ten": k[1], "l1": a, "l2": b, "delta": round(b-a, 2) if a is not None else None}
            if b < 5: ph.append(item)
            elif 5 <= b < 6.5: qu.append(item)
            if b >= 8: bo.append(item)
            if a is not None:
                if b-a >= 1.5: ti.append(item)
                if b-a <= -2: th.append(item)
        ph.sort(key=lambda x: x["l2"]); qu.sort(key=lambda x: x["l2"]); bo.sort(key=lambda x: -x["l2"])
        ti.sort(key=lambda x: -x["delta"]); th.sort(key=lambda x: x["delta"])
        class_lists[lop][s] = {"phudao": ph, "quantam": qu, "boiduong": bo, "tienbo": ti, "thutlui": th}

block_lists = {}
for s in SUBJECTS:
    ph, bd = [], []
    for k in all_keys:
        b = d2.get(k, {}).get(s) if k in d2 else None
        if b is None: continue
        a = d1.get(k, {}).get(s) if k in d1 else None
        item = {"lop": k[0], "ten": k[1], "l1": a, "l2": b, "delta": round(b-a, 2) if a is not None else None}
        if b < 5: ph.append(item)
        if b >= 8: bd.append(item)
    ph.sort(key=lambda x: x["l2"]); bd.sort(key=lambda x: -x["l2"])
    block_lists[s] = {"phudao": ph, "boiduong": bd}

top_changes = {}
for s in SUBJECTS:
    arr = [{"lop": k[0], "ten": k[1], "l1": d1[k][s], "l2": d2[k][s], "delta": round(d2[k][s]-d1[k][s], 2)}
           for k in set(d1) & set(d2) if d1[k].get(s) is not None and d2[k].get(s) is not None]
    top_changes[s] = {"giam_manh_nhat": sorted(arr, key=lambda x: x["delta"])[:10],
                      "tang_manh_nhat": sorted(arr, key=lambda x: -x["delta"])[:10]}

data["meta"]["cap_nhat"] = datetime.now().strftime("%H:%M %d/%m/%Y")
for e in log:
    if e not in data["meta"]["nguon"] and (e.startswith("12A") or "+" in e or "~" in e[:3]):
        pass
data["meta"]["nguon"] = data["meta"]["nguon"] + [f"batch 11/09 23:09: {os.path.basename(f)}" for f in newfiles]
data["meta"]["siso_l1"] = len(d1); data["meta"]["siso_l2"] = len(d2)
data["subjects_stats"] = subjects_stats
data["overview_classes"] = [{"lop": lop,
    "siso_l1": sum(1 for k in all_keys if k[0] == lop and k in d1),
    "siso_l2": sum(1 for k in all_keys if k[0] == lop and k in d2),
    "stats": class_stats[lop]} for lop in classes]
data["class_stats"] = class_stats; data["class_lists"] = class_lists
data["block_lists"] = block_lists; data["top_changes"] = top_changes; data["students"] = students
json.dump(data, open(OUT, "w", encoding="utf-8"), ensure_ascii=False)
print("wrote", OUT, os.path.getsize(OUT))
for s in SUBJECTS:
    st = subjects_stats[s]
    print(f"{s}: n1={st['n1']} TB1={st['tb1']} | n2={st['n2']} TB2={st['tb2']} | Δ={st['delta']} | <5: {st['duoi5_l1']}→{st['duoi5_l2']} | ≥8 L2: {st['gioi_l2']}")
